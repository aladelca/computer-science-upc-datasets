from __future__ import annotations

import json
import tempfile
import textwrap
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from typing import Sequence

from PIL import Image, ImageDraw, ImageFont

DEFAULT_SLIDE_WIDTH = 1920
DEFAULT_SLIDE_HEIGHT = 1080
DEFAULT_ACCENT = "#d64045"
TEXT = "#111111"
MUTED = "#666666"
LIGHT_PANEL = "#f7f3ee"
SOFT_BACKGROUND = "#fcfbf8"
DIVIDER = "#d9d9d9"


class SlideKind(StrEnum):
    COVER = "cover"
    COMIC = "comic"
    CHART = "chart"
    DIVIDER = "divider"


@dataclass
class RenderedSlide:
    index: int
    kind: SlideKind
    title: str
    subtitle: str | None
    prompt_context: str | None
    image: Image.Image
    notes: str | None = None


@dataclass(frozen=True)
class WeekMaterialPaths:
    week_number: int
    topic_slug: str
    markdown_path: Path
    notebook_path: Path | None
    generator_script_path: Path
    presentation_pdf: Path
    presentation_pptx: Path
    class_script_path: Path


@dataclass(frozen=True)
class MarkdownSection:
    heading: str
    bullets: list[str]


@dataclass(frozen=True)
class OutlineFonts:
    week: ImageFont.FreeTypeFont
    title: ImageFont.FreeTypeFont
    heading: ImageFont.FreeTypeFont
    body: ImageFont.FreeTypeFont
    small: ImageFont.FreeTypeFont
    tiny: ImageFont.FreeTypeFont


def _load_font(size: int, *, bold: bool = False) -> ImageFont.FreeTypeFont:
    from matplotlib import font_manager

    weight = "bold" if bold else "normal"
    font_path = font_manager.findfont(
        font_manager.FontProperties(family="DejaVu Sans", weight=weight),
        fallback_to_default=True,
    )
    return ImageFont.truetype(font_path, size=size)


def load_outline_fonts() -> OutlineFonts:
    return OutlineFonts(
        week=_load_font(42, bold=True),
        title=_load_font(72, bold=True),
        heading=_load_font(48, bold=True),
        body=_load_font(30),
        small=_load_font(24),
        tiny=_load_font(18),
    )


def new_slide(
    *,
    width: int = DEFAULT_SLIDE_WIDTH,
    height: int = DEFAULT_SLIDE_HEIGHT,
    background: str = "white",
) -> Image.Image:
    return Image.new("RGBA", (width, height), background)


def _text_size(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.FreeTypeFont,
) -> tuple[int, int]:
    if not text:
        return 0, 0
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=10)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text(
    draw: ImageDraw.ImageDraw,
    text: str,
    font: ImageFont.FreeTypeFont,
    max_width: int,
) -> str:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        cleaned = paragraph.strip()
        if not cleaned:
            lines.append("")
            continue
        words = cleaned.split()
        current = words[0]
        for word in words[1:]:
            candidate = f"{current} {word}"
            if draw.textlength(candidate, font=font) <= max_width:
                current = candidate
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return "\n".join(lines)


def _draw_title_block(
    slide: Image.Image,
    fonts: OutlineFonts,
    *,
    week_label: str,
    title: str,
    subtitle: str | None,
    accent: str,
) -> None:
    draw = ImageDraw.Draw(slide)
    draw.rounded_rectangle(
        (120, 140, 1800, 940),
        radius=44,
        fill=SOFT_BACKGROUND,
        outline=DIVIDER,
        width=4,
    )
    draw.rounded_rectangle((120, 140, 220, 940), radius=44, fill=accent)
    draw.text((280, 230), week_label, fill=accent, font=fonts.week)
    draw.text((280, 330), title, fill=TEXT, font=fonts.title)
    if subtitle:
        wrapped = wrap_text(draw, subtitle, fonts.body, 1320)
        draw.multiline_text(
            (280, 470),
            wrapped,
            fill=TEXT,
            font=fonts.body,
            spacing=12,
        )


def _draw_section_block(
    slide: Image.Image,
    fonts: OutlineFonts,
    *,
    heading: str,
    bullets: Sequence[str],
    footer: str | None,
    accent: str,
) -> None:
    draw = ImageDraw.Draw(slide)
    draw.rounded_rectangle(
        (110, 90, 1810, 980),
        radius=40,
        fill="white",
        outline=DIVIDER,
        width=4,
    )
    draw.rounded_rectangle((110, 90, 1810, 220), radius=40, fill=LIGHT_PANEL)
    draw.rectangle((110, 175, 1810, 220), fill=LIGHT_PANEL)
    draw.text((170, 132), heading, fill=TEXT, font=fonts.heading)
    draw.rounded_rectangle((150, 280, 1770, 840), radius=28, fill=SOFT_BACKGROUND)
    cursor = 330
    for bullet in bullets:
        wrapped = wrap_text(draw, f"• {bullet}", fonts.body, 1480)
        draw.multiline_text(
            (210, cursor),
            wrapped,
            fill=TEXT,
            font=fonts.body,
            spacing=10,
        )
        _, height = _text_size(draw, wrapped, fonts.body)
        cursor += height + 26
    if footer:
        draw.text((170, 890), footer, fill=accent, font=fonts.small)


def _presentation_stem(week_number: int) -> str:
    return f"week_{week_number}"


def _script_stem(week_number: int) -> str:
    return f"generate_week_{week_number}_presentation.py"


def _class_script_name(week_number: int, topic_slug: str) -> str:
    return f"week_{week_number:02d}_{topic_slug}_class_script.md"


def discover_week_material_paths(
    repo_root: Path,
    week_number: int,
) -> WeekMaterialPaths:
    content_dir = repo_root / "big_data_course_content"
    notebook_dir = content_dir / "notebooks"
    markdown_candidates = sorted(
        path
        for path in content_dir.glob(f"week_{week_number:02d}_*.md")
        if not path.name.endswith("_class_script.md")
    )
    if not markdown_candidates:
        raise FileNotFoundError(f"No markdown content found for week {week_number:02d}")
    markdown_path = markdown_candidates[0]
    topic_slug = markdown_path.stem.removeprefix(f"week_{week_number:02d}_")

    notebook_path = next(
        iter(sorted(notebook_dir.glob(f"week_{week_number:02d}_*.ipynb"))),
        None,
    )
    presentation_stem = _presentation_stem(week_number)
    return WeekMaterialPaths(
        week_number=week_number,
        topic_slug=topic_slug,
        markdown_path=markdown_path,
        notebook_path=notebook_path,
        generator_script_path=repo_root / "scripts" / _script_stem(week_number),
        presentation_pdf=content_dir / "presentations" / f"{presentation_stem}.pdf",
        presentation_pptx=content_dir / "presentations" / f"{presentation_stem}.pptx",
        class_script_path=content_dir / _class_script_name(week_number, topic_slug),
    )


def _clean_markdown_line(line: str) -> str:
    cleaned = line.strip()
    for marker in ("`", "**", "__"):
        cleaned = cleaned.replace(marker, "")
    if cleaned.startswith("> "):
        cleaned = cleaned[2:]
    if cleaned.startswith("- "):
        cleaned = cleaned[2:]
    if cleaned[:2].isdigit() and cleaned[2:4] == ". ":
        cleaned = cleaned[4:]
    if cleaned.startswith("### "):
        cleaned = cleaned[4:]
    if cleaned.startswith("## "):
        cleaned = cleaned[3:]
    if cleaned.startswith("# "):
        cleaned = cleaned[2:]
    return " ".join(cleaned.split())


def parse_markdown_outline(
    markdown_path: Path,
    *,
    max_sections: int = 8,
    max_bullets_per_section: int = 4,
) -> tuple[str, list[str], list[MarkdownSection]]:
    title = markdown_path.stem.replace("_", " ").title()
    intro_lines: list[str] = []
    sections: list[MarkdownSection] = []
    current_heading: str | None = None
    current_lines: list[str] = []
    in_code_block = False

    for raw_line in markdown_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            in_code_block = not in_code_block
            continue
        if in_code_block or not stripped:
            continue
        if stripped.startswith("$$") or stripped.startswith(r"\["):
            continue
        if stripped.startswith("# "):
            title = _clean_markdown_line(stripped)
            continue
        if stripped.startswith("## "):
            if current_heading:
                sections.append(
                    MarkdownSection(
                        heading=current_heading,
                        bullets=current_lines[:max_bullets_per_section],
                    )
                )
            current_heading = _clean_markdown_line(stripped)
            current_lines = []
            continue
        cleaned = _clean_markdown_line(stripped)
        if not cleaned:
            continue
        if current_heading is None:
            if len(intro_lines) < 3:
                intro_lines.append(cleaned)
            continue
        if len(current_lines) < max_bullets_per_section:
            current_lines.append(cleaned)

    if current_heading and len(sections) < max_sections:
        sections.append(
            MarkdownSection(
                heading=current_heading,
                bullets=current_lines[:max_bullets_per_section],
            )
        )

    return title, intro_lines, sections[:max_sections]


def parse_notebook_outline(
    notebook_path: Path | None,
    *,
    max_items: int = 6,
) -> list[str]:
    if notebook_path is None or not notebook_path.exists():
        return []
    payload = json.loads(notebook_path.read_text(encoding="utf-8"))
    headings: list[str] = []
    for cell in payload.get("cells", []):
        if cell.get("cell_type") != "markdown":
            continue
        for line in cell.get("source", []):
            stripped = line.strip()
            if stripped.startswith("#"):
                cleaned = _clean_markdown_line(stripped)
                if cleaned and cleaned not in headings:
                    headings.append(cleaned)
                break
        if len(headings) >= max_items:
            break
    return headings[:max_items]


def build_outline_scaffold_slides(
    paths: WeekMaterialPaths,
    *,
    accent: str = DEFAULT_ACCENT,
) -> list[RenderedSlide]:
    fonts = load_outline_fonts()
    title, intro_lines, sections = parse_markdown_outline(paths.markdown_path)
    notebook_headings = parse_notebook_outline(paths.notebook_path)

    slides: list[RenderedSlide] = []

    def push(
        image: Image.Image,
        *,
        kind: SlideKind,
        title: str,
        subtitle: str | None = None,
        prompt_context: str | None = None,
        notes: str | None = None,
    ) -> None:
        slides.append(
            RenderedSlide(
                index=len(slides) + 1,
                kind=kind,
                title=title,
                subtitle=subtitle,
                prompt_context=prompt_context,
                image=image,
                notes=notes,
            )
        )

    cover = new_slide(background="white")
    _draw_title_block(
        cover,
        fonts,
        week_label=f"Week {paths.week_number}",
        title=title,
        subtitle=" ".join(intro_lines[:2]) if intro_lines else None,
        accent=accent,
    )
    push(
        cover,
        kind=SlideKind.COVER,
        title=title,
        subtitle=" ".join(intro_lines[:2]) if intro_lines else None,
        prompt_context="Scaffold cover slide generated from week markdown.",
    )

    if intro_lines:
        intro_slide = new_slide(background="white")
        _draw_section_block(
            intro_slide,
            fonts,
            heading="Opening Frame",
            bullets=intro_lines,
            footer=f"Source: {paths.markdown_path.name}",
            accent=accent,
        )
        push(
            intro_slide,
            kind=SlideKind.CHART,
            title="Opening Frame",
            subtitle="Summary extracted from the week markdown.",
            prompt_context="Introductory scaffold slide generated from markdown preamble.",
        )

    for section in sections:
        slide = new_slide(background="white")
        bullets = section.bullets or ["Replace this scaffold bullet set with a week-specific explanation."]
        _draw_section_block(
            slide,
            fonts,
            heading=section.heading,
            bullets=bullets,
            footer="Scaffold slide. Replace with custom visuals and tighter pacing.",
            accent=accent,
        )
        push(
            slide,
            kind=SlideKind.CHART,
            title=section.heading,
            subtitle="Scaffolded from markdown structure.",
            prompt_context=f"Section scaffold slide for {section.heading}.",
        )

    if notebook_headings:
        notebook_slide = new_slide(background="white")
        _draw_section_block(
            notebook_slide,
            fonts,
            heading="Notebook Flow",
            bullets=notebook_headings,
            footer=f"Notebook: {paths.notebook_path.name if paths.notebook_path else 'N/A'}",
            accent=accent,
        )
        push(
            notebook_slide,
            kind=SlideKind.CHART,
            title="Notebook Flow",
            subtitle="Practice checkpoints detected from the week notebook.",
            prompt_context="Notebook checkpoints scaffold slide.",
        )

    closing = new_slide(background="white")
    _draw_section_block(
        closing,
        fonts,
        heading="Next Steps",
        bullets=[
            "Replace outline slides with week-specific visuals and mathematical diagrams.",
            "Review formulas, examples, and narrative transitions against the markdown and notebook.",
            "After approval, wire in the redesign pipeline for AI beautification if needed.",
        ],
        footer="This deck is a deterministic starting point, not the final teaching version.",
        accent=accent,
    )
    push(
        closing,
        kind=SlideKind.DIVIDER,
        title="Next Steps",
        subtitle="Scaffold closing slide.",
        prompt_context="Closing scaffold slide with next-step checklist.",
    )

    return slides


def _coerce_slide_images(
    slides: Sequence[RenderedSlide | Image.Image],
) -> list[Image.Image]:
    images: list[Image.Image] = []
    for slide in slides:
        if isinstance(slide, Image.Image):
            images.append(slide)
            continue
        images.append(slide.image)
    return images


def save_pdf(
    slides: Sequence[RenderedSlide | Image.Image],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    rgb_slides = [image.convert("RGB") for image in _coerce_slide_images(slides)]
    first, *rest = rgb_slides
    first.save(output_path, save_all=True, append_images=rest, resolution=150.0)


def save_pptx(
    slides: Sequence[RenderedSlide | Image.Image],
    output_path: Path,
) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_path = Path(temp_dir)
        for index, slide_image in enumerate(_coerce_slide_images(slides), 1):
            slide = presentation.slides.add_slide(blank_layout)
            image_path = temp_path / f"slide_{index:02d}.png"
            slide_image.convert("RGB").save(image_path)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )

        if len(presentation.slides) > len(slides):
            xml_slides = presentation.slides._sldIdLst  # type: ignore[attr-defined]
            xml_slides.remove(xml_slides[0])

        presentation.save(output_path)


def build_generator_scaffold(paths: WeekMaterialPaths) -> str:
    return textwrap.dedent(
        f'''\
        from __future__ import annotations

        import argparse
        import sys
        from pathlib import Path

        ROOT = Path(__file__).resolve().parents[1]
        SRC = ROOT / "src"
        if str(SRC) not in sys.path:
            sys.path.insert(0, str(SRC))

        from upc_datasets.weekly_presentation import (  # noqa: E402
            build_outline_scaffold_slides,
            discover_week_material_paths,
            save_pdf,
            save_pptx,
        )

        WEEK_NUMBER = {paths.week_number}
        WEEK_PATHS = discover_week_material_paths(ROOT, week_number=WEEK_NUMBER)


        def parse_args() -> argparse.Namespace:
            parser = argparse.ArgumentParser(
                description="Generate the Week {paths.week_number} scaffold deck."
            )
            parser.add_argument(
                "--output",
                type=Path,
                default=WEEK_PATHS.presentation_pdf,
                help="Output PDF path.",
            )
            parser.add_argument(
                "--output-pptx",
                type=Path,
                default=WEEK_PATHS.presentation_pptx,
                help="Output PowerPoint path.",
            )
            return parser.parse_args()


        def main() -> None:
            args = parse_args()
            slides = build_outline_scaffold_slides(WEEK_PATHS)
            save_pdf(slides, args.output)
            if str(args.output_pptx).strip():
                save_pptx(slides, args.output_pptx)
                print(
                    f"Generated {{len(slides)}} scaffold slides at {{args.output}} "
                    f"and {{args.output_pptx}}"
                )
            else:
                print(f"Generated {{len(slides)}} scaffold slides at {{args.output}}")


        if __name__ == "__main__":
            main()

        # Next steps:
        # 1. Replace outline scaffold slides with week-specific visuals.
        # 2. Add mathematical diagrams and examples from the week markdown and notebook.
        # 3. After user approval, wire in a redesign profile and AI beautification pipeline.
        '''
    )


def write_generator_scaffold(
    paths: WeekMaterialPaths,
    output_path: Path | None = None,
    *,
    force: bool = False,
) -> Path:
    target = output_path or paths.generator_script_path
    if target.exists() and not force:
        raise FileExistsError(f"Refusing to overwrite existing scaffold: {target}")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(build_generator_scaffold(paths), encoding="utf-8")
    return target
